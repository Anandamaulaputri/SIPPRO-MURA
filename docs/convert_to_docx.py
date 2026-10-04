import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import re

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="B0B0B0", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_footer_with_page_number(doc, font_name="Times New Roman"):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    
    # Add top border line to footer
    p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="4" w:color="CBD5E1"/></w:pBdr>')
    p._p.get_or_add_pPr().append(p_border)

    run_doc = p.add_run("SIPPRO MURA — Dokumen Kebutuhan Sistem (PRD)  |  Halaman ")
    run_doc.font.name = font_name
    run_doc.font.size = Pt(9.5)
    run_doc.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Add XML PAGE field
    fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    p._p.append(fldSimple)

def build_docx(md_path, docx_path):
    doc = Document()

    # Standard Academic Margins (Normal 2.54 cm / 1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    FONT_NAME = 'Times New Roman'
    BODY_SIZE = Pt(12)
    LINE_SPACING = 1.5

    # Configure Normal style
    style_normal = doc.styles['Normal']
    style_normal.font.name = FONT_NAME
    style_normal.font.size = BODY_SIZE
    style_normal.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    style_normal.paragraph_format.line_spacing = LINE_SPACING
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    add_footer_with_page_number(doc, FONT_NAME)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    in_table = False
    table_lines = []

    def flush_table(t_lines):
        if not t_lines:
            return
        rows_data = []
        for l in t_lines:
            parts = [c.strip() for c in l.strip().strip('|').split('|')]
            if any(re.match(r'^:?-+:?$', p) for p in parts):
                continue
            rows_data.append(parts)

        if not rows_data:
            return

        num_cols = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color="CBD5E1")

        for r_idx, row in enumerate(rows_data):
            is_header = (r_idx == 0)
            for c_idx in range(num_cols):
                cell = table.cell(r_idx, c_idx)
                text = row[c_idx] if c_idx < len(row) else ""
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15  # Tables stay compact for clean layout

                clean_text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
                clean_text = re.sub(r'\*(.*?)\*', r'\1', clean_text)
                clean_text = clean_text.replace('<br>', '\n')

                run = p.add_run(clean_text)
                run.font.name = FONT_NAME
                
                if is_header:
                    set_cell_background(cell, "1E3A8A") # Navy header
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    run.font.size = Pt(11)
                else:
                    if r_idx % 2 == 1:
                        set_cell_background(cell, "F8FAFC")
                    else:
                        set_cell_background(cell, "FFFFFF")
                    run.font.size = Pt(10.5)
                    if text.startswith('**') and text.endswith('**'):
                        run.bold = True

        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_before = Pt(0)
        p_spacer.paragraph_format.space_after = Pt(4)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Code block
        if stripped.startswith('```'):
            in_code_block = not in_code_block
            i += 1
            continue

        if in_code_block:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(line.rstrip('\n'))
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            i += 1
            continue

        # Tables
        if stripped.startswith('|') and '|' in stripped[1:]:
            in_table = True
            table_lines.append(stripped)
            i += 1
            continue
        else:
            if in_table:
                flush_table(table_lines)
                table_lines = []
                in_table = False

        if not stripped:
            i += 1
            continue

        # Horizontal Rule
        if stripped in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="94A3B8"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            i += 1
            continue

        # Callouts / Quotes (> Catatan Penting:, etc.)
        if stripped.startswith('>'):
            alert_text = stripped.lstrip('> ').strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = LINE_SPACING
            
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="12" w:color="1E3A8A"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            
            clean_text = re.sub(r'\[!(.*?)\]', r'[\1]', alert_text)
            clean_text = re.sub(r'\*\*(.*?)\*\*', r'\1', clean_text)
            run = p.add_run(clean_text)
            run.font.name = FONT_NAME
            run.font.size = Pt(11)
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            i += 1
            continue

        # Headings: Times New Roman, bold
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.2
            run = p.add_run(stripped[2:])
            run.font.name = FONT_NAME
            run.font.size = Pt(15)
            run.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            i += 1
            continue

        if stripped.startswith('## '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.2
            run = p.add_run(stripped[3:])
            run.font.name = FONT_NAME
            run.font.size = Pt(13.5)
            run.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            i += 1
            continue

        if stripped.startswith('### '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.2
            run = p.add_run(stripped[4:])
            run.font.name = FONT_NAME
            run.font.size = Pt(12.5)
            run.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            i += 1
            continue

        if stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.2
            run = p.add_run(stripped[5:])
            run.font.name = FONT_NAME
            run.font.size = Pt(12)
            run.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            i += 1
            continue

        # Bullet lists (- or *)
        if stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = LINE_SPACING
            content = stripped[2:]
            _add_formatted_text(p, content, FONT_NAME, BODY_SIZE)
            i += 1
            continue

        # Numbered list (1. / 2.)
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = LINE_SPACING
            _add_formatted_text(p, num_match.group(2), FONT_NAME, BODY_SIZE)
            i += 1
            continue

        # Indented sub-items in Daftar Isi (e.g. 1.1 Latar Belakang)
        if re.match(r'^\s+(\d+\.\d+)\s+(.*)$', line):
            sub_match = re.match(r'^\s+(\d+\.\d+)\s+(.*)$', line)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.2
            _add_formatted_text(p, f"{sub_match.group(1)} {sub_match.group(2)}", FONT_NAME, Pt(11.5))
            i += 1
            continue

        # Regular paragraph: Times New Roman 12 pt, Spacing 1.5, Justify (Rata Kiri-Kanan)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = LINE_SPACING
        _add_formatted_text(p, stripped, FONT_NAME, BODY_SIZE)
        i += 1

    if in_table:
        flush_table(table_lines)

    doc.save(docx_path)
    print(f"Document successfully updated with TNR 12pt, 1.5 spacing, Justify: {docx_path}")

def _add_formatted_text(paragraph, text, font_name, font_size):
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            r = paragraph.add_run(tok[2:-2])
            r.font.name = font_name
            r.font.size = font_size
            r.bold = True
        elif tok.startswith('*') and tok.endswith('*'):
            r = paragraph.add_run(tok[1:-1])
            r.font.name = font_name
            r.font.size = font_size
            r.italic = True
        else:
            r = paragraph.add_run(tok)
            r.font.name = font_name
            r.font.size = font_size

if __name__ == '__main__':
    import os
    base_dir = os.path.dirname(os.path.abspath(__file__))
    md_file = os.path.join(base_dir, 'PRD_SIPPRO.md')
    docx_file = os.path.join(base_dir, 'PRD_SIPPRO_MURA.docx')
    build_docx(md_file, docx_file)
